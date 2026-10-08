"""Text extraction from uploaded files (PDF, DOCX, TXT).

Each format has its own small function so a new format (e.g. ODT, HTML) can be
added without touching the others. All failures are converted into
`DocumentError`s with a message the user can act on.
"""

from __future__ import annotations

import io
from pathlib import Path

from app.core.errors import DocumentError, UnsupportedFormatError
from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(filename: str, content: bytes, max_bytes: int | None = None) -> str:
    """Return the raw text of an uploaded file, dispatching on its extension."""
    if not content:
        raise DocumentError(f"The file '{filename}' is empty.")
    if max_bytes is not None and len(content) > max_bytes:
        raise DocumentError(
            f"The file '{filename}' is too large ({len(content) / 1_048_576:.1f} MB). "
            f"Maximum allowed size is {max_bytes / 1_048_576:.0f} MB."
        )

    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"Unsupported file type '{extension or 'unknown'}'. Please upload a PDF, DOCX or TXT file."
        )

    if extension == ".pdf":
        text = _extract_pdf(filename, content)
    elif extension == ".docx":
        text = _extract_docx(filename, content)
    else:
        text = _decode_text(content)

    logger.info("Extracted %d characters from %s file", len(text), extension)
    return text


def _extract_pdf(filename: str, content: bytes) -> str:
    # Check the magic number first: a renamed file would otherwise produce a
    # confusing low-level parser error.
    if not content.lstrip()[:5].startswith(b"%PDF"):
        raise DocumentError(f"'{filename}' does not look like a valid PDF file.")
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentError(f"'{filename}' is password-protected. Please upload an unprotected PDF.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        logger.warning("PDF parsing failed: %s", type(exc).__name__)
        raise DocumentError(f"Could not read the PDF '{filename}'. The file may be corrupted.") from exc

    text = "\n\n".join(pages)
    if not text.strip():
        # Typical for scanned CVs: the PDF only contains images. OCR is out of
        # scope for the MVP, so tell the user explicitly instead of silently
        # continuing with an empty document.
        raise DocumentError(
            f"No text could be extracted from '{filename}'. It may be a scanned image; "
            "please upload a text-based PDF, a DOCX file or paste the text."
        )
    return text


def _extract_docx(filename: str, content: bytes) -> str:
    if not content.startswith(b"PK"):  # DOCX files are ZIP archives
        raise DocumentError(f"'{filename}' does not look like a valid DOCX file.")
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # python-docx raises a variety of zip/XML errors
        logger.warning("DOCX parsing failed: %s", type(exc).__name__)
        raise DocumentError(f"Could not read the DOCX '{filename}'. The file may be corrupted.") from exc

    lines = [p.text for p in document.paragraphs]
    # Many CV templates put skills or dates in tables; ignoring them would lose
    # exactly the information we need for matching.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat text
    return "\n".join(lines)


def _decode_text(content: bytes) -> str:
    # UTF-8 first (with or without BOM); cp1252 is the usual culprit for text
    # files saved on Windows. latin-1 never fails, so it is the final fallback.
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("Could not decode the text file.")  # pragma: no cover
