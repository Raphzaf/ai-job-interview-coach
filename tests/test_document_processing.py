"""Text extraction, cleaning and validation."""

import io

import pytest

from app.core.errors import DocumentError, UnsupportedFormatError
from app.document_processing.cleaner import clean_text, normalize_for_matching, validate_document
from app.document_processing.extractors import extract_text


# --- cleaning -------------------------------------------------------------------

def test_clean_text_normalises_whitespace_and_bullets():
    raw = "SKILLS\r\n\r\n\r\n•  Python,\tSQL\n● Docker  and  AWS\n\n\n\nEnd"
    assert clean_text(raw) == "SKILLS\n\n- Python, SQL\n- Docker and AWS\n\nEnd"


def test_clean_text_preserves_meaningful_tokens():
    raw = "C++, C#, Node.js, 5+ years, M.Sc. in AI, AWS-SAA certification"
    assert clean_text(raw) == raw  # nothing meaningful may be stripped


def test_clean_text_fixes_pdf_artefacts():
    # Ligature, zero-width space and hyphenated line break from PDF extraction.
    assert clean_text("ﬁnance devel-\nopment​ team") == "finance development team"


def test_clean_text_keeps_date_ranges_and_capitalised_hyphens():
    assert "2019-\n2021" in clean_text("2019-\n2021")
    assert "Front-\nEnd" in clean_text("Front-\nEnd")


def test_normalize_for_matching_is_lowercase_single_line():
    assert normalize_for_matching("Python\n  SQL") == "python sql"


# --- validation -----------------------------------------------------------------

def test_validate_rejects_empty():
    with pytest.raises(DocumentError, match="empty"):
        validate_document("   ", "CV", 80, 1000)


def test_validate_rejects_too_short():
    with pytest.raises(DocumentError, match="too short"):
        validate_document("Python developer", "CV", 80, 1000)


def test_validate_rejects_too_long():
    with pytest.raises(DocumentError, match="too long"):
        validate_document("word " * 500, "CV", 10, 1000)


def test_validate_rejects_non_text():
    with pytest.raises(DocumentError, match="readable text"):
        validate_document("%$#@ 1234 5678 9012 !!!! ???? " * 5, "CV", 10, 10_000)


def test_validate_accepts_normal_document():
    text = "Experienced data scientist with Python and SQL skills. " * 3
    assert validate_document(text, "CV", 80, 10_000) == text.strip()


def test_processor_prepare_cleans_and_validates(processor, sample_cv_text):
    doc = processor.prepare(sample_cv_text, "cv", "cv.txt")
    assert doc.doc_type == "cv" and "Brightleaf" in doc.text
    with pytest.raises(DocumentError):
        processor.prepare("", "job", "job.txt")


# --- extraction -----------------------------------------------------------------

def test_extract_txt_with_utf8_and_cp1252():
    assert extract_text("cv.txt", "Résumé – data".encode("utf-8")) == "Résumé – data"
    assert extract_text("cv.txt", "Résumé".encode("cp1252")) == "Résumé"


def test_extract_docx_includes_tables():
    import docx

    document = docx.Document()
    document.add_paragraph("Alex Morgan - Machine Learning Engineer")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, FAISS"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("cv.docx", buffer.getvalue())
    assert "Alex Morgan" in text and "Skills | Python, FAISS" in text


def _make_pdf(text: str | None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    if text:
        pdf.drawString(72, 800, text)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def test_extract_pdf():
    assert "Machine Learning Engineer" in extract_text("cv.pdf", _make_pdf("Alex Morgan - Machine Learning Engineer"))


def test_extract_pdf_without_text_layer_raises():
    with pytest.raises(DocumentError, match="scanned"):
        extract_text("scan.pdf", _make_pdf(None))


@pytest.mark.parametrize(
    "filename, content, error, message",
    [
        ("cv.exe", b"MZ...", UnsupportedFormatError, "Unsupported"),
        ("cv.txt", b"", DocumentError, "empty"),
        ("cv.pdf", b"this is not a pdf", DocumentError, "valid PDF"),
        ("cv.pdf", b"%PDF-1.4 garbage garbage", DocumentError, "Could not read"),
        ("cv.docx", b"not a zip", DocumentError, "valid DOCX"),
        ("cv.docx", b"PK\x03\x04 broken zip", DocumentError, "Could not read"),
    ],
)
def test_extract_errors(filename, content, error, message):
    with pytest.raises(error, match=message):
        extract_text(filename, content)


def test_extract_rejects_oversized_files():
    with pytest.raises(DocumentError, match="too large"):
        extract_text("cv.txt", b"a" * 2_000, max_bytes=1_000)
