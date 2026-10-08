"""Generate PDF and DOCX versions of the fictional sample CV.

    python scripts/make_sample_documents.py

The TXT file in data/examples is the single source of truth; this script lets
the demo show that PDF and DOCX uploads go through the same pipeline.
Requires the dev dependencies (reportlab).
"""

from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

EXAMPLES = Path(__file__).resolve().parents[1] / "data" / "examples"


def is_heading(line: str) -> bool:
    return line.isupper() and len(line.split()) <= 4


def build_pdf(lines: list[str], target: Path) -> None:
    styles = getSampleStyleSheet()
    story = []
    for line in lines:
        if not line.strip():
            story.append(Spacer(1, 6))
        elif is_heading(line):
            story.append(Paragraph(escape(line), styles["Heading3"]))
        elif line.startswith("- "):
            story.append(Paragraph("• " + escape(line[2:]), styles["BodyText"]))
        else:
            story.append(Paragraph(escape(line), styles["BodyText"]))
    SimpleDocTemplate(str(target), pagesize=A4, title="Sample CV (fictional)").build(story)


def build_docx(lines: list[str], target: Path) -> None:
    document = docx.Document()
    for line in lines:
        if not line.strip():
            continue
        if is_heading(line):
            document.add_heading(line.title(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(str(target))


if __name__ == "__main__":
    source_lines = (EXAMPLES / "sample_cv.txt").read_text(encoding="utf-8").splitlines()
    build_pdf(source_lines, EXAMPLES / "sample_cv.pdf")
    build_docx(source_lines, EXAMPLES / "sample_cv.docx")
    print("Wrote sample_cv.pdf and sample_cv.docx to", EXAMPLES)
